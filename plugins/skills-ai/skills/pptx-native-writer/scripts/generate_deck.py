#!/usr/bin/env python3
"""generate_deck.py — build a NATIVE, editable .pptx in the Stoic design system.

Requires: python-pptx  (pip install python-pptx)
Part of the StoicDesign `pptx-native-writer` skill.

Usage:
  python3 generate_deck.py --spec deck.json --brand ../assets/brand_guide.json --out deck.pptx

See the skill SKILL.md for the deck.json schema. Slide types: title, section, content, stat, close.
"""
import argparse
import json
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
except ImportError:
    sys.exit("python-pptx not installed. Run:  pip install python-pptx")


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def load_json(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def textbox(slide, text, left, top, width, height, *, size, font, color,
            bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.name = font
    f.bold = bold
    f.color.rgb = rgb(color)
    return tb


def fill_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(color)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True)
    ap.add_argument("--brand", required=True)
    ap.add_argument("--out", default="deck.pptx")
    a = ap.parse_args()

    spec = load_json(a.spec)
    brand = load_json(a.brand)
    C, F, S = brand["colors"], brand["fonts"], brand["sizes_pt"]
    lesson = spec.get("lesson", "")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    MX = 0.9
    total = len(spec["slides"])

    for i, s in enumerate(spec["slides"], 1):
        t = s.get("type", "content")
        slide = prs.slides.add_slide(blank)
        dark = t in ("title", "section", "stat", "close")
        fill_bg(slide, C["navy"] if dark else C["parchment"])
        ink = C["cream"] if dark else C["ink"]
        muted = C["cream_muted"] if dark else C["ink_muted"]
        accent = C["data_gold"] if dark else C["wax"]

        if s.get("kicker"):
            textbox(slide, s["kicker"].upper(), MX, 0.6, 11, 0.4,
                    size=S["kicker"], font=F["ui"], color=accent, bold=True)

        if t == "title":
            textbox(slide, s.get("title", ""), MX, 2.2, 11.0, 2.6,
                    size=S["display"], font=F["display"], color=ink, bold=True)
            if s.get("subtitle"):
                textbox(slide, s["subtitle"], MX, 5.1, 9.0, 1.2,
                        size=S["body"], font=F["ui"], color=muted)
        elif t == "section":
            textbox(slide, s.get("title", ""), MX, 2.8, 11.5, 1.8,
                    size=S["title"] + 16, font=F["display"], color=ink, bold=True)
        elif t == "stat":
            textbox(slide, s.get("stat", ""), MX, 2.1, 11.5, 2.2,
                    size=S["stat"], font=F["mono"], color=accent, bold=True)
            if s.get("caption"):
                textbox(slide, s["caption"], MX, 4.8, 10.0, 1.2,
                        size=S["body"], font=F["ui"], color=muted)
        elif t == "close":
            textbox(slide, s.get("title", ""), MX, 2.4, 11.5, 2.0,
                    size=S["title"], font=F["display"], color=ink, bold=True)
            if s.get("subtitle"):
                textbox(slide, s["subtitle"], MX, 4.6, 10.5, 1.2,
                        size=S["body"], font=F["ui"], color=muted)
            if s.get("sources"):
                textbox(slide, s["sources"], MX, 6.3, 11.5, 0.6,
                        size=S["source"], font=F["mono"], color=muted)
        else:  # content
            textbox(slide, s.get("title", ""), MX, 1.3, 11.5, 1.6,
                    size=S["title"], font=F["display"], color=ink, bold=True)
            bullets = s.get("bullets", [])
            if bullets:
                tb = slide.shapes.add_textbox(Inches(MX), Inches(3.1), Inches(11.5), Inches(3.4))
                tf = tb.text_frame
                tf.word_wrap = True
                for j, b in enumerate(bullets):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    run = p.add_run()
                    run.text = b
                    f = run.font
                    f.size = Pt(S["body"] + 2)
                    f.name = F["ui"]
                    f.color.rgb = rgb(ink)
                    p.space_after = Pt(10)

        # footer spine
        rule = C["cream_muted"] if dark else C["ink_muted"]
        textbox(slide, f"{brand.get('footer', 'STOIC EDU')} · {lesson}", MX, 7.0, 9.5, 0.4,
                size=S["footer"], font=F["ui"], color=rule)
        textbox(slide, f"{i:02d} / {total:02d}", 11.0, 7.0, 1.4, 0.4,
                size=S["footer"], font=F["mono"], color=rule, align=PP_ALIGN.RIGHT)

    prs.save(a.out)
    print(f"native pptx -> {a.out}  ({total} slides)")


if __name__ == "__main__":
    main()
